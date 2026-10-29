"""Extract plain text from ~15 file formats.

Every reader is optional: if the backing library is missing we skip the file
with a clear reason instead of crashing the whole ingestion run.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

TEXT_EXT = {".txt", ".md", ".markdown", ".rst", ".org", ".log", ".csv", ".tsv", ".json",
            ".yaml", ".yml", ".xml", ".tex", ".sql", ".ini", ".toml", ".cfg", ".conf"}
CODE_EXT = {".py", ".js", ".ts", ".jsx", ".tsx", ".java", ".go", ".rs", ".c", ".h", ".cpp",
            ".hpp", ".cs", ".rb", ".php", ".sh", ".bash", ".zsh", ".swift", ".kt", ".scala",
            ".r", ".m", ".mm", ".lua", ".pl", ".pm", ".vim", ".el", ".clj", ".ex", ".exs"}

# Directories never worth walking into.
SKIP_DIRS = {".git", ".svn", ".hg", "node_modules", "__pycache__", ".venv", "venv",
             "env", ".tox", "dist", "build", ".idea", ".vscode", ".engram", ".mypy_cache",
             ".pytest_cache", ".gradle", "target", "vendor", ".next", ".cache"}

MAX_BYTES = 20 * 1024 * 1024  # skip >20 MB files


@dataclass
class Document:
    text: str
    source: str
    title: str = ""
    meta: dict = field(default_factory=dict)

    @property
    def ok(self) -> bool:
        return bool(self.text and self.text.strip())


def _read_text_file(p: Path) -> str:
    raw = p.read_bytes()
    for enc in ("utf-8", "utf-8-sig", "gb18030", "big5", "latin-1"):
        try:
            return raw.decode(enc)
        except (UnicodeDecodeError, LookupError):
            continue
    return raw.decode("utf-8", errors="replace")


def _pdf(p: Path) -> str:
    from pypdf import PdfReader
    r = PdfReader(str(p))
    return "\n".join((pg.extract_text() or "") for pg in r.pages)


def _docx(p: Path) -> str:
    import docx
    d = docx.Document(str(p))
    parts = [para.text for para in d.paragraphs]
    for tbl in d.tables:
        for row in tbl.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(parts)


def _xlsx(p: Path) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(str(p), read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"# sheet: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                out.append(" | ".join(cells))
    wb.close()
    return "\n".join(out)


def _epub(p: Path) -> str:
    import ebooklib
    from ebooklib import epub
    from bs4 import BeautifulSoup
    book = epub.read_epub(str(p), options={"ignore_ncx": True})
    out = []
    for item in book.get_items():
        if item.get_type() == ebooklib.ITEM_DOCUMENT:
            soup = BeautifulSoup(item.get_content(), "html.parser")
            out.append(soup.get_text(separator="\n"))
    return "\n".join(out)


def _html(p: Path) -> str:
    raw = _read_text_file(p)
    try:
        from bs4 import BeautifulSoup
        soup = BeautifulSoup(raw, "html.parser")
        for tag in soup(["script", "style", "nav", "footer", "header"]):
            tag.decompose()
        return soup.get_text(separator="\n")
    except ImportError:
        import re
        return re.sub(r"<[^>]+>", " ", raw)


def _notebook(p: Path) -> str:
    import json
    nb = json.loads(_read_text_file(p))
    out = []
    for cell in nb.get("cells", []):
        src = cell.get("source", [])
        out.append("".join(src) if isinstance(src, list) else str(src))
    return "\n".join(out)


READERS = {
    ".pdf": _pdf, ".docx": _docx, ".xlsx": _xlsx, ".xlsm": _xlsx,
    ".epub": _epub, ".html": _html, ".htm": _html, ".ipynb": _notebook,
}


def parse_file(path: str | Path, with_meta: bool = True) -> Document:
    """Parse one file into a Document. Never raises — inspect .ok / .meta['error']."""
    p = Path(path)
    doc = Document(text="", source=str(p), title=p.name)
    try:
        if not p.exists() or not p.is_file():
            doc.meta["error"] = "not a file"
            return doc
        if p.stat().st_size > MAX_BYTES:
            doc.meta["error"] = "file too large (>20 MB)"
            return doc
        ext = p.suffix.lower()
        if ext in READERS:
            doc.text = READERS[ext](p)
            doc.meta["reader"] = ext.lstrip(".")
        elif ext in TEXT_EXT or ext in CODE_EXT or ext == "":
            doc.text = _read_text_file(p)
            doc.meta["reader"] = "text"
        else:
            doc.meta["error"] = f"unsupported extension {ext}"
            return doc
        doc.text = doc.text.strip()
        if with_meta:
            st = p.stat()
            doc.meta.update(size=st.st_size, mtime=st.st_mtime, ext=ext)
            doc.title = _guess_title(doc.text, p.name)
    except ImportError as e:
        doc.meta["error"] = f"missing optional dependency: {e.name or e}"
    except Exception as e:  # noqa: BLE001 - ingestion must never die on one bad file
        doc.meta["error"] = f"{type(e).__name__}: {e}"
    return doc


def _guess_title(text: str, fallback: str) -> str:
    for line in text.splitlines()[:12]:
        s = line.strip().lstrip("#").strip()
        if 2 <= len(s) <= 120:
            return s
    return fallback


def walk(path: str | Path) -> list[Path]:
    """Collect ingestable files from a file or directory tree."""
    p = Path(path)
    if p.is_file():
        return [p]
    out = []
    for f in p.rglob("*"):
        if f.is_dir():
            continue
        if any(part in SKIP_DIRS for part in f.parts):
            continue
        if f.name.startswith("."):
            continue
        out.append(f)
    return sorted(out)
